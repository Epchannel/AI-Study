"""End-to-end academy contracts on isolated SQLite, with offline generation."""
import json
import asyncio
import uuid
from datetime import datetime, timedelta, timezone, date
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest
import pytest_asyncio
from fastapi import FastAPI
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from sqlalchemy.pool import StaticPool

from database import Base, get_db
from routers import academy
from services.academy import build_schedule, streak_summary, parse_generation, extract_segments


@pytest_asyncio.fixture
async def portal(tmp_path, monkeypatch):
    from sse_starlette.sse import AppStatus
    monkeypatch.setattr(AppStatus, "should_exit_event", None)
    engine = create_async_engine("sqlite+aiosqlite:///" + (tmp_path / "academy.db").as_posix())
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    app = FastAPI()
    app.include_router(academy.router)

    async def db():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_db] = db
    monkeypatch.setattr(academy, "async_session", factory)
    monkeypatch.setattr(academy.settings, "upload_dir", str(tmp_path))
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        yield client
    await engine.dispose()


async def register(client, email):
    response = await client.post("/api/academy/register", json={"name": email.split("@")[0], "email": email, "password": "test-only-12345"})
    assert response.status_code == 201, response.text
    body = response.json()
    return body["user"], {"Authorization": "Bearer " + body["access_token"]}


async def classroom(client, teacher):
    response = await client.post("/api/academy/classes", headers=teacher, json={"name": "Lớp Toán 10", "subject": "Toán", "grade": "10"})
    assert response.status_code == 201, response.text
    return response.json()


@pytest.mark.asyncio
async def test_attendance_authorization_and_upsert(portal):
    _, ta = await register(portal, "attendance-teacher@test.vn")
    student, sa = await register(portal, "attendance-student@test.vn")
    outsider, oa = await register(portal, "attendance-outsider@test.vn")
    room = await classroom(portal, ta)
    await portal.post("/api/academy/classes/join", headers=sa, json={"code": room["join_code"]})
    endpoint = f"/api/academy/classes/{room['id']}/attendance?day=2026-10-30"
    update = f"/api/academy/classes/{room['id']}/students/{student['id']}/attendance"
    body = {"day": "2026-10-30", "status": "present", "note": ""}
    assert (await portal.get(endpoint, headers=ta)).json() == []
    for auth in (sa, oa):
        assert (await portal.get(endpoint, headers=auth)).status_code == 403
        assert (await portal.put(update, headers=auth, json=body)).status_code == 403
    assert (await portal.put(update.replace(student['id'], outsider['id']), headers=ta, json=body)).status_code == 404
    assert (await portal.put(update, headers=ta, json=body)).status_code == 200
    body.update(status="late", note="Đến muộn 10 phút")
    assert (await portal.put(update, headers=ta, json=body)).status_code == 200
    rows = (await portal.get(endpoint, headers=ta)).json()
    assert len(rows) == 1
    assert rows[0] == {"user_id": student["id"], "name": student["name"], **body}
    assert (await portal.put(update, headers=ta, json={**body, "status": "unknown"})).status_code == 422


async def material(client, auth, class_id, parent=None):
    response = await client.post(f"/api/academy/classes/{class_id}/materials", headers=auth,
        files={"file": ("notes.md", "# Phương trình\nNghiệm của x + 2 = 5 là x = 3.".encode())},
        data={"chapter": "Chương 1", "lesson": "Phương trình", "parent_id": parent or ""})
    assert response.status_code == 201, response.text
    return response.json()


async def publish(client, auth, doc):
    response = await client.put(f"/api/academy/materials/{doc['id']}", headers=auth,
        json={"title": doc["title"], "status": "published", **doc["data"]})
    assert response.status_code == 200, response.text


@pytest.mark.asyncio
async def test_account_roles_and_class_isolation(portal):
    teacher, ta = await register(portal, "teacher@test.vn")
    student, sa = await register(portal, "student@test.vn")
    _, outsider = await register(portal, "outsider@test.vn")
    assert teacher["role"] == "teacher" and student["role"] == "student"
    assert (await portal.get("/api/academy/me")).status_code == 401
    denied = await portal.post("/api/academy/classes", headers=sa, json={"name": "No", "subject": "Toán", "grade": "10"})
    assert denied.status_code == 403
    room = await classroom(portal, ta)
    assert (await portal.get(f"/api/academy/classes/{room['id']}/materials", headers=outsider)).status_code == 403
    assert (await portal.post("/api/academy/classes/join", headers=sa, json={"code": room["join_code"]})).status_code == 200
    rooms = (await portal.get("/api/academy/classes", headers=sa)).json()
    assert "join_code" not in rooms[0]
    assert (await portal.get(f"/api/academy/classes/{room['id']}/report", headers=sa)).status_code == 403
    # Profile payload cannot promote a student to teacher.
    assert (await portal.put("/api/academy/profile", headers=sa, json={"role": "teacher"})).status_code == 200
    assert (await portal.get("/api/academy/me", headers=sa)).json()["role"] == "student"


@pytest.mark.asyncio
async def test_material_draft_publish_version_and_download(portal):
    _, ta = await register(portal, "teacher@test.vn")
    _, sa = await register(portal, "student@test.vn")
    _, other = await register(portal, "other@test.vn")
    room = await classroom(portal, ta)
    await portal.post("/api/academy/classes/join", headers=sa, json={"code": room["join_code"]})
    doc = await material(portal, ta, room["id"])
    assert doc["status"] == "draft"
    assert "path" not in doc
    assert (await portal.get(f"/api/academy/classes/{room['id']}/materials", headers=sa)).json() == []
    assert (await portal.get(f"/api/academy/materials/{doc['id']}/download", headers=sa)).status_code == 404
    await publish(portal, ta, doc)
    assert (await portal.get(f"/api/academy/materials/{doc['id']}/download", headers=sa)).status_code == 200
    assert (await portal.get(f"/api/academy/materials/{doc['id']}/download", headers=other)).status_code == 403
    new_doc = await material(portal, ta, room["id"], doc["id"])
    assert new_doc["data"]["version"] == 2
    await publish(portal, ta, new_doc)
    visible = (await portal.get(f"/api/academy/classes/{room['id']}/materials", headers=sa)).json()
    assert [d["id"] for d in visible] == [new_doc["id"]]
    denied = await portal.post(f"/api/academy/classes/{room['id']}/materials", headers=sa, files={"file": ("a.txt", b"test")})
    assert denied.status_code == 403


QUIZ = {"items": [{"question": "2 + 2 bằng bao nhiêu?", "options": ["1", "2", "3", "4"], "answer": 3,
                   "explanation": "Cộng hai với hai được bốn.", "topic": "Phép cộng", "difficulty": "foundation"}]}


@pytest.mark.asyncio
async def test_teacher_overview_aggregates_only_owned_classes(portal):
    from models.academy import LearningAsset, AcademyAttempt
    teacher, ta = await register(portal, "overview-teacher@test.vn")
    student, sa = await register(portal, "overview-student@test.vn")
    _, oa = await register(portal, "overview-outsider@test.vn")
    room = await classroom(portal, ta)
    await portal.post("/api/academy/classes/join", headers=sa, json={"code": room["join_code"]})
    foreign = await classroom(portal, ta)
    await portal.post("/api/academy/classes/join", headers=sa, json={"code": foreign["join_code"]})
    doc = await material(portal, ta, room["id"])
    await publish(portal, ta, doc)
    async with academy.async_session() as db:
        asset = LearningAsset(user_id=uuid.UUID(teacher["id"]), classroom_id=uuid.UUID(room["id"]), kind="quiz", status="draft", title="Duyệt", data={"items": [], "subject": "Toán"})
        db.add(asset)
        await db.flush()
        db.add(AcademyAttempt(user_id=uuid.UUID(student["id"]), asset_id=asset.id, data={"score": 2, "total": 3}, created_at=datetime.now(timezone.utc)))
        await db.commit()
    overview = await portal.get("/api/academy/teacher/overview", headers=ta)
    assert overview.status_code == 200, overview.text
    assert overview.json()["class_count"] == 2
    assert overview.json()["student_count"] == 1
    assert overview.json()["materials"] == {"total": 1, "draft": 0, "published": 1}
    assert overview.json()["assignments"]["pending_review"] == 1
    assert overview.json()["active_students"] == 1
    assert overview.json()["recent"][0]["student_name"] == "overview-student" and overview.json()["recent"][0]["asset_title"] == "Duyệt"
    assert (await portal.get("/api/academy/teacher/overview", headers=sa)).status_code == 403
    export = await portal.get(f"/api/academy/classes/{room['id']}/report.csv", headers=ta)
    assert export.status_code == 200 and "Học sinh" in export.text and "Bài giao" in export.text
    assert (await portal.get(f"/api/academy/classes/{foreign['id']}/report.csv", headers=oa)).status_code == 403
    analytics = await portal.get(f"/api/academy/assets/{asset.id}/analytics", headers=ta)
    assert analytics.status_code == 200 and analytics.json()["attempt_count"] == 1
    assert (await portal.get(f"/api/academy/assets/{asset.id}/analytics", headers=sa)).status_code == 403


@pytest.mark.asyncio
async def test_new_assignment_due_date_cannot_be_past(portal, monkeypatch):
    import services.llm.router as llm
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(chat=AsyncMock(return_value=(json.dumps({"items": QUIZ["items"] * 3}), {}))))
    _, auth = await register(portal, "due-date@test.vn")
    yesterday = (datetime.now(timezone(timedelta(hours=7))) - timedelta(days=1)).date().isoformat()
    response = await portal.post("/api/academy/assets/generate", headers=auth, json={"kind": "quiz", "topic": "Toán", "count": 3, "due_date": yesterday})
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_selected_chat_document_isolated_persisted_and_rechecked(portal, monkeypatch):
    import services.llm.router as llm
    prompts = []
    async def stream(system, message):
        prompts.append(system)
        yield "Giải thích từ tài liệu [1]."
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(stream_chat=stream))
    _, teacher = await register(portal, "source-teacher@test.vn")
    _, student = await register(portal, "source-student@test.vn")
    room = await classroom(portal, teacher)
    await portal.post("/api/academy/classes/join", headers=student, json={"code": room["join_code"]})
    first = await material(portal, teacher, room["id"])
    second = await material(portal, teacher, room["id"])
    second["data"]["segments"] = [{"label": "Khác", "text": "UNSELECTED_SOURCE_MARKER"}]
    await publish(portal, teacher, first)
    await publish(portal, teacher, second)
    body = {"subject": "Toán", "classroom_id": room["id"], "material_id": first["id"]}
    created = await portal.post("/api/academy/threads", headers=student, json=body)
    assert created.status_code == 201, created.text
    thread = created.json()
    response = await portal.post(f"/api/academy/threads/{thread['id']}/chat", headers=student, json={"message": "Giải thích bài học"})
    assert "event: done" in response.text
    assert "UNSELECTED_SOURCE_MARKER" not in prompts[0] and "x + 2 = 5" in prompts[0]
    saved = (await portal.get(f"/api/academy/threads/{thread['id']}", headers=student)).json()
    assert saved["data"]["material_id"] == first["id"]
    assert {s["material_id"] for s in saved["data"]["messages"][-1]["sources"]} == {first["id"]}
    generator = AsyncMock(return_value=(json.dumps({"items": QUIZ["items"] * 3}), {}))
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(stream_chat=stream, chat=generator))
    exercise = await portal.post(f"/api/academy/threads/{thread['id']}/chat", headers=student, json={"message": "Tạo 3 câu quiz từ tài liệu"})
    assert "event: action" in exercise.text and "event: done" in exercise.text
    generator.assert_awaited_once()
    assert "UNSELECTED_SOURCE_MARKER" not in generator.call_args.args[0]
    saved = (await portal.get(f"/api/academy/threads/{thread['id']}", headers=student)).json()
    asset = (await portal.get(f"/api/academy/assets/{saved['data']['messages'][-1]['asset_id']}", headers=student)).json()
    assert {s["material_id"] for s in asset["data"]["sources"]} == {first["id"]}
    await portal.put(f"/api/academy/materials/{first['id']}", headers=teacher, json={"title": first["title"], **first["data"], "status": "archived"})
    denied = await portal.post(f"/api/academy/threads/{thread['id']}/chat", headers=student, json={"message": "Hỏi tiếp"})
    assert denied.status_code == 404 and len(prompts) == 1
    assert uuid.UUID(thread["id"]) not in academy._chat_locks
    other_room = await classroom(portal, teacher)
    assert (await portal.post("/api/academy/threads", headers=teacher, json={**body, "classroom_id": other_room["id"]})).status_code == 404
    assert (await portal.post("/api/academy/threads", headers=student, json={"material_id": second["id"]})).status_code == 422


@pytest.mark.asyncio
async def test_schedule_reschedule_duration_restore_ownership_and_completion(portal):
    from models.academy import AcademySchedule
    student, auth = await register(portal, "schedule@test.vn")
    _, other = await register(portal, "other-schedule@test.vn")
    day = (datetime.now(timezone(timedelta(hours=7))) + timedelta(days=1)).date().isoformat()
    past = (datetime.now(timezone(timedelta(hours=7))) - timedelta(days=1)).date().isoformat()
    async with academy.async_session() as db:
        entries = [AcademySchedule(user_id=uuid.UUID(student["id"]), status=status,
                   data={"title": "Ôn Toán", "subject": "Toán", "start": f"{d}T{start}:00+07:00", "end": f"{d}T{end}:00+07:00", "minutes": 30, "goal": "Ôn tập"})
                   for d, start, end, status in [(day, "18:00", "18:30", "planned"), (day, "18:00", "18:30", "skipped"), (past, "18:00", "18:30", "planned")]]
        db.add_all(entries)
        await db.commit()
        first, skipped, old = [str(entry.id) for entry in entries]
    async def patch(id, body, headers=auth):
        return await portal.patch(f"/api/academy/schedule/{id}", headers=headers, json=body)
    assert (await patch(first, {"status": "skipped"}, other)).status_code == 404
    assert (await patch(first, {"status": "completed"})).status_code == 409
    assert (await patch(skipped, {"status": "planned"})).status_code == 409
    assert (await patch(first, {"status": "planned", "start": f"{day}T19:00:00"})).status_code == 422
    moved = await patch(first, {"status": "planned", "start": f"{day}T12:00:00Z", "end": f"{day}T13:15:00Z"})
    assert moved.status_code == 200, moved.text
    assert moved.json()["data"]["minutes"] == 75
    assert moved.json()["data"]["start"] == f"{day}T19:00:00+07:00"
    assert (await patch(skipped, {"status": "planned"})).status_code == 200
    assert (await patch(old, {"status": "completed"})).status_code == 200
    assert (await patch(old, {"status": "completed"})).status_code == 200
    progress = (await portal.get("/api/academy/progress", headers=auth)).json()
    assert progress["completed_activities"] == 1 and progress["streak"]["current"] == 1
    assert (await patch(old, {"status": "planned"})).status_code == 409
    assert (await patch(first, {"status": "planned", "start": f"{day}T23:30:00+07:00", "end": f"{day}T23:30:30+07:00"})).status_code == 422


@pytest.mark.asyncio
async def test_teacher_report_excludes_private_and_other_class_work_and_orders_attempts(portal):
    from models.academy import LearningAsset, AcademyAttempt, AcademyCard
    _, teacher = await register(portal, "report-teacher@test.vn")
    student, auth = await register(portal, "report-student@test.vn")
    room = await classroom(portal, teacher)
    second = await classroom(portal, teacher)
    for target in [room, second]:
        await portal.post("/api/academy/classes/join", headers=auth, json={"code": target["join_code"]})
    async with academy.async_session() as db:
        assets = [LearningAsset(user_id=uuid.UUID(student["id"]), classroom_id=class_id, title=title, status="published", kind="quiz", data={"items": []})
                  for class_id, title in [(uuid.UUID(room["id"]), "Bài của lớp"), (uuid.UUID(second["id"]), "Bài lớp khác"), (None, "Bài học riêng")]]
        db.add_all(assets)
        await db.flush()
        now = datetime.now(timezone.utc)
        for asset in assets:
            # Insert newer first to ensure report uses timestamps rather than insertion order.
            for days, score in [(0, 1), (1, 0)]:
                db.add(AcademyAttempt(user_id=uuid.UUID(student["id"]), asset_id=asset.id, created_at=now - timedelta(days=days), data={"score": score, "total": 1}))
            db.add(AcademyCard(user_id=uuid.UUID(student["id"]), asset_id=asset.id, card_index=0, data={"reps": 1}))
        await db.commit()
        class_asset = str(assets[0].id)
    response = await portal.get(f"/api/academy/classes/{room['id']}/report", headers=teacher)
    assert response.status_code == 200, response.text
    report = response.json()
    assert len(report["assignments"]) == 1
    assert report["assignments"][0]["submitted"] == 1 and report["assignments"][0]["completion_rate"] == 100
    learner = report["students"][0]
    assert learner["cards_reviewed"] == 1
    assert len(learner["attempts"]) == 2
    assert all(attempt["asset_id"] == class_asset for attempt in learner["attempts"])
    assert learner["attempts"][-1]["data"]["score"] == 1
    assert set(learner) == {"id", "name", "attempts", "cards_reviewed", "last_activity", "average_score"}
    assert learner["average_score"] == 50 and learner["last_activity"]


@pytest.mark.asyncio
async def test_exercise_uses_selected_material_version_and_rejects_other_class(portal, monkeypatch):
    import services.llm.router as llm
    generator = AsyncMock(return_value=(json.dumps({"items": QUIZ["items"] * 3}), {}))
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(chat=generator))
    _, auth = await register(portal, "material-source@test.vn")
    room = await classroom(portal, auth)
    original = await material(portal, auth, room["id"])
    revision = await material(portal, auth, room["id"], original["id"])
    payload = {"kind": "quiz", "topic": "Phương trình", "count": 3, "classroom_id": room["id"], "material_id": revision["id"]}
    result = await portal.post("/api/academy/assets/generate", headers=auth, json=payload)
    assert result.status_code == 201, result.text
    asset = (await portal.get(f"/api/academy/assets/{result.json()['id']}", headers=auth)).json()
    assert asset["status"] == "draft"
    assert asset["data"]["sources"] == [{"material_id": revision["id"], "title": revision["title"], "version": 2}]
    assert "x + 2 = 5" in generator.call_args.args[0]
    other_room = await classroom(portal, auth)
    payload["classroom_id"] = other_room["id"]
    assert (await portal.post("/api/academy/assets/generate", headers=auth, json=payload)).status_code == 404
    generator.assert_awaited_once()


@pytest.mark.asyncio
async def test_concurrent_learning_writes_preserve_mastery_and_schedule(portal, monkeypatch):
    import services.llm.router as llm
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(chat=AsyncMock(return_value=(json.dumps({"items": QUIZ["items"] * 5}), {}))))
    _, auth = await register(portal, "concurrent@test.vn")
    generated = await portal.post("/api/academy/assets/generate", headers=auth, json={"kind": "quiz", "topic": "Phép cộng"})
    assert generated.status_code == 201, generated.text
    aid = generated.json()["id"]
    await portal.put(f"/api/academy/assets/{aid}", headers=auth, json={"title": "Quiz", "status": "published", "items": QUIZ["items"]})
    responses = await asyncio.gather(*[
        portal.post(f"/api/academy/assets/{aid}/attempts", headers=auth, json={"answers": [3], "submission_id": str(uuid.uuid4())})
        for _ in range(3)
    ])
    assert all(r.status_code == 200 for r in responses), [r.text for r in responses]
    progress = (await portal.get("/api/academy/progress", headers=auth)).json()
    assert progress["mastery"]["Toán: Phép cộng"] == {"correct": 3, "total": 3}
    start = datetime.now(timezone(timedelta(hours=7))).date() + timedelta(days=1)
    profile = {"grade": "10", "level": "foundation", "subjects": ["Toán"], "goal": "Ôn tập",
               "availability": [{"weekday": start.weekday(), "start": "18:00", "end": "21:00"}]}
    assert (await portal.put("/api/academy/profile", headers=auth, json=profile)).status_code == 200
    payload = {"start": start.isoformat(), "end": start.isoformat(), "subjects": ["Toán"], "daily_minutes": 60}
    responses = await asyncio.gather(*[portal.post("/api/academy/schedule/generate", headers=auth, json=payload) for _ in range(2)])
    assert all(r.status_code == 200 for r in responses), [r.text for r in responses]
    entries = (await portal.get("/api/academy/schedule", headers=auth)).json()
    assert entries and sum(r["data"]["minutes"] for r in entries) <= 60
    ordered = sorted(entries, key=lambda r: r["data"]["start"])
    assert all(a["data"]["end"] <= b["data"]["start"] for a, b in zip(ordered, ordered[1:]))


@pytest.mark.asyncio
async def test_quiz_approval_answer_secrecy_scoring_and_streak(portal, monkeypatch):
    import services.llm.router as llm
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(chat=AsyncMock(return_value=(json.dumps({"items": QUIZ["items"] * 5}), {}))))
    _, ta = await register(portal, "teacher@test.vn")
    _, sa = await register(portal, "student@test.vn")
    room = await classroom(portal, ta)
    await portal.post("/api/academy/classes/join", headers=sa, json={"code": room["join_code"]})
    generated = await portal.post("/api/academy/assets/generate", headers=ta,
        json={"kind": "quiz", "topic": "Phép cộng", "classroom_id": room["id"]})
    assert generated.status_code == 201, generated.text
    aid = generated.json()["id"]
    assert (await portal.get(f"/api/academy/assets/{aid}", headers=sa)).status_code == 404
    assert (await portal.put(f"/api/academy/assets/{aid}", headers=ta, json={"title": "Phép cộng", "status": "published", "items": QUIZ["items"]})).status_code == 200
    student_view = (await portal.get(f"/api/academy/assets/{aid}", headers=sa)).json()
    assert "answer" not in student_view["data"]["items"][0]
    assert "explanation" not in student_view["data"]["items"][0]
    import uuid
    body = {"answers": [0], "submission_id": str(uuid.uuid4())}
    first = await portal.post(f"/api/academy/assets/{aid}/attempts", headers=sa, json=body)
    assert first.status_code == 200, first.text
    assert first.json()["data"]["score"] == 0
    assert first.json()["data"]["results"][0]["answer"] == 3
    await portal.post(f"/api/academy/assets/{aid}/attempts", headers=sa, json=body)
    progress = (await portal.get("/api/academy/progress", headers=sa)).json()
    assert len(progress["attempts"]) == 1 and progress["streak"]["current"] == 1
    assert progress["mastery"]["Toán: Phép cộng"] == {"correct": 0, "total": 1}
    report = (await portal.get(f"/api/academy/classes/{room['id']}/report", headers=ta)).json()
    assert len(report["students"][0]["attempts"]) == 1


@pytest.mark.asyncio
async def test_flashcard_fsrs_private_state_and_duplicate_review(portal, monkeypatch):
    import services.llm.router as llm
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(chat=AsyncMock(return_value=(json.dumps({"items": [{"front": "2+2", "back": "4", "topic": "Cộng"}] * 5}), {}))))
    _, ta = await register(portal, "teacher@test.vn")
    _, sa = await register(portal, "student@test.vn")
    r = await portal.post("/api/academy/assets/generate", headers=sa, json={"kind": "flashcards", "topic": "Cộng"})
    aid = r.json()["id"]
    assert (await portal.get(f"/api/academy/assets/{aid}", headers=ta)).status_code == 404
    reviewed = await portal.post(f"/api/academy/assets/{aid}/cards/0/review", headers=sa, json={"rating": 3})
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["reps"] == 1 and reviewed.json()["due"]
    duplicate = await portal.post(f"/api/academy/assets/{aid}/cards/0/review", headers=sa, json={"rating": 3})
    assert duplicate.status_code == 409
    assert (await portal.get("/api/academy/progress", headers=sa)).json()["completed_activities"] == 1


@pytest.mark.asyncio
async def test_conversation_persistence_and_ownership(portal, monkeypatch):
    import services.llm.router as llm
    async def stream(*args):
        yield "Chào em. "
        yield "Mình học Toán nhé!"
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(stream_chat=stream))
    _, ta = await register(portal, "teacher@test.vn")
    _, sa = await register(portal, "student@test.vn")
    thread = (await portal.post("/api/academy/threads", headers=sa, json={"subject": "Toán"})).json()
    tid = thread["id"]
    assert (await portal.get(f"/api/academy/threads/{tid}", headers=ta)).status_code == 404
    r = await portal.post(f"/api/academy/threads/{tid}/chat", headers=sa, json={"message": "Dạy em phép cộng"})
    assert "event: done" in r.text
    saved = (await portal.get(f"/api/academy/threads/{tid}", headers=sa)).json()
    assert len(saved["data"]["messages"]) == 2
    assert saved["data"]["messages"][1]["content"] == "Chào em. Mình học Toán nhé!"
    # Simply opening/chatting does not grant a study streak.
    assert (await portal.get("/api/academy/progress", headers=sa)).json()["streak"]["current"] == 0


def test_vietnam_day_boundary_and_longest_streak():
    now = datetime(2026, 10, 29, 18, tzinfo=timezone.utc)  # Oct 30 in Vietnam
    events = [SimpleNamespace(created_at=datetime(2026, 10, day, 18, tzinfo=timezone.utc)) for day in (26, 27, 28)]
    summary = streak_summary(events, now)
    assert summary["current"] == 3 and summary["longest"] == 3 and not summary["today_completed"]
    assert summary["days"] == ["2026-10-27", "2026-10-28", "2026-10-29"]
    assert streak_summary(events, now + timedelta(days=1))["current"] == 0


def test_schedule_respects_availability_daily_limit_and_existing_work():
    start = date(2026, 11, 2)
    profile = {"availability": [{"weekday": 0, "start": "19:00", "end": "21:00"}], "goal": "Ôn thi"}
    existing = [{"start": "2026-11-02T19:00:00+07:00", "end": "2026-11-02T19:30:00+07:00"}]
    result = build_schedule(profile, start, start, ["Toán", "Tiếng Anh"], 60, existing)
    assert sum(r["minutes"] for r in result) <= 60
    assert result[0]["start"] == "2026-11-02T19:40:00+07:00"
    assert all(r["end"] <= "2026-11-02T21:00:00+07:00" for r in result)
    assert sum(r["minutes"] for r in result) == 30  # 30 minutes already booked
    again = build_schedule(profile, start, start, ["Toán"], 60, existing + result)
    assert again == []  # Rebuilding the schedule cannot exceed the daily study budget.


def test_invalid_ai_output_not_published_and_text_extraction(tmp_path):
    from fastapi import HTTPException
    with pytest.raises(HTTPException):
        parse_generation('{"items":[{"question":"Q"}]}', "quiz", 5)
    path = tmp_path / "notes.md"
    path.write_text("# Bài học\nNội dung tiếng Việt.", encoding="utf-8")
    assert "tiếng Việt" in extract_segments(path)[0]["text"]


@pytest.mark.asyncio
async def test_chat_creates_saved_exercise_with_terminal_action(portal, monkeypatch):
    import services.llm.router as llm
    generator = AsyncMock(return_value=(json.dumps({"items": QUIZ["items"] * 3}), {}))
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(chat=generator))
    _, auth = await register(portal, "teacher@test.vn")
    thread = (await portal.post("/api/academy/threads", headers=auth, json={"subject": "Toán"})).json()
    response = await portal.post(f"/api/academy/threads/{thread['id']}/chat", headers=auth,
                                 json={"message": "Tạo 3 câu quiz về phép cộng"})
    assert "event: action" in response.text and "event: done" in response.text, response.text
    saved = (await portal.get(f"/api/academy/threads/{thread['id']}", headers=auth)).json()
    asset_id = saved["data"]["messages"][-1]["asset_id"]
    asset = (await portal.get(f"/api/academy/assets/{asset_id}", headers=auth)).json()
    assert len(asset["data"]["items"]) == 3
    assert thread["id"] not in {str(t) for t in academy._chat_locks}
    generator.assert_awaited_once()


@pytest.mark.asyncio
async def test_failed_stream_is_not_saved_and_releases_thread(portal, monkeypatch):
    import services.llm.router as llm
    async def broken(*args):
        yield "Nội dung dở dang"
        raise ConnectionError("Upstream disconnected")
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(stream_chat=broken))
    _, auth = await register(portal, "teacher@test.vn")
    thread = (await portal.post("/api/academy/threads", headers=auth, json={})).json()
    response = await portal.post(f"/api/academy/threads/{thread['id']}/chat", headers=auth,
                                 json={"message": "Dạy em phép cộng"})
    assert "event: error" in response.text and "event: done" not in response.text
    saved = (await portal.get(f"/api/academy/threads/{thread['id']}", headers=auth)).json()
    assert saved["data"]["messages"] == []
    assert uuid.UUID(thread["id"]) not in academy._chat_locks


@pytest.mark.asyncio
async def test_browser_disconnect_cancels_upstream_without_saving_partial_chat(portal, monkeypatch):
    import services.llm.router as llm
    closed = asyncio.Event()
    disconnected = asyncio.Event()
    async def slow_stream(*args):
        try:
            yield "Phần trả lời đầu tiên"
            await asyncio.Event().wait()
        finally:
            closed.set()
    monkeypatch.setattr(llm, "get_llm_client", lambda: SimpleNamespace(stream_chat=slow_stream))
    _, auth = await register(portal, "disconnect@test.vn")
    thread = (await portal.post("/api/academy/threads", headers=auth, json={})).json()
    payload = json.dumps({"message": "Giải thích phép cộng"}).encode()
    delivered = False
    output = []
    async def receive():
        nonlocal delivered
        if not delivered:
            delivered = True
            return {"type": "http.request", "body": payload, "more_body": False}
        await disconnected.wait()
        return {"type": "http.disconnect"}
    async def send(message):
        output.append(message)
        if message["type"] == "http.response.body" and b"event: content" in message.get("body", b""):
            disconnected.set()
    scope = {"type": "http", "asgi": {"version": "3.0", "spec_version": "2.3"},
             "http_version": "1.1", "method": "POST", "scheme": "http",
             "path": f"/api/academy/threads/{thread['id']}/chat", "raw_path": b"/", "query_string": b"",
             "root_path": "", "server": ("test", 80), "client": ("127.0.0.1", 1234),
             "headers": [(b"content-type", b"application/json"), (b"authorization", auth["Authorization"].encode())]}
    await asyncio.wait_for(portal._transport.app(scope, receive, send), timeout=5)
    assert closed.is_set() and disconnected.is_set()
    assert not any(b"event: done" in message.get("body", b"") for message in output)
    assert uuid.UUID(thread["id"]) not in academy._chat_locks
    saved = (await portal.get(f"/api/academy/threads/{thread['id']}", headers=auth)).json()
    assert saved["data"]["messages"] == []


@pytest.mark.parametrize("extension", ["docx", "pptx", "pdf"])
def test_office_pdf_extraction_preserves_source_positions(tmp_path, extension):
    path = tmp_path / ("lesson." + extension)
    if extension == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph("Linear equations: x + 2 = 5.")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "x = 3"
        document.save(path)
    elif extension == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Linear equations: x + 2 = 5."
        presentation.save(path)
    else:
        # A small valid PDF with a text stream; no renderer or external fixture needed.
        objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
                   b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                   b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
                   b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
        content = b"BT /F1 12 Tf 72 720 Td (Linear equations: x + 2 = 5.) Tj ET"
        objects.append(b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream")
        pdf = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for i, obj in enumerate(objects, 1):
            offsets.append(len(pdf))
            pdf.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
        xref = len(pdf)
        pdf.extend(b"xref\n0 6\n0000000000 65535 f \n")
        for offset in offsets[1:]:
            pdf.extend(f"{offset:010d} 00000 n \n".encode())
        pdf.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
        path.write_bytes(pdf)
    result = extract_segments(path)
    assert "x + 2 = 5" in result[0]["text"]
    assert result[0]["label"] == {"docx": "Nội dung", "pptx": "Slide 1", "pdf": "Trang 1"}[extension]
